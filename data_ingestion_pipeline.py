from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Dict, List, Optional, Union
from urllib.parse import urlparse

PathLike = Union[str, Path]


class LoadData:
    def __init__(self, request_timeout: int = 30) -> None:
        self.request_timeout = request_timeout
        self.supported_formats = {
            ".pdf": self.load_pdf,
            ".ppt": self.load_ppt,
            ".pptx": self.load_ppt,
            ".docx": self.load_word,
            ".doc": self.load_word,
            ".txt": self.load_text,
            ".xlsx": self.load_excel,
            ".xls": self.load_excel,
            ".mp3": self.load_audio,
            ".wav": self.load_audio,
            ".mp4": self.load_video,
            ".avi": self.load_video,
            ".mov": self.load_video,
            ".url": self.load_url_file,
        }

    def load_pdf(self, data: PathLike, pages: Optional[List[int]] = None) -> str:
        from unstructured.partition.pdf import partition_pdf

        elements = partition_pdf(filename=str(data), strategy="ocr_only", pages=pages)
        return "\n".join(element.text for element in elements if element.text)

    def load_ppt(self, data: PathLike) -> str:
        from pptx import Presentation

        path = Path(data)
        if path.suffix.lower() == ".ppt":
            raise ValueError(
                "Legacy .ppt files must be converted to .pptx before processing."
            )

        presentation = Presentation(str(path))
        return "\n".join(
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        )

    def load_word(self, data: PathLike) -> str:
        path = Path(data)
        if path.suffix.lower() == ".doc":
            import aspose.words as aw

            with TemporaryDirectory() as temp_dir:
                converted_path = Path(temp_dir) / f"{path.stem}.docx"
                aw.Document(str(path)).save(str(converted_path))
                return self._read_docx(converted_path)
        return self._read_docx(path)

    @staticmethod
    def _read_docx(path: Path) -> str:
        from docx import Document

        document = Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    @staticmethod
    def load_excel(data: PathLike, sheet_name: Optional[str] = None) -> str:
        import pandas as pd

        workbook = pd.read_excel(data, sheet_name=sheet_name)
        if isinstance(workbook, dict):
            return "\n\n".join(
                f"{name}\n{frame.to_string(index=False)}"
                for name, frame in workbook.items()
            )
        return workbook.to_string(index=False)

    @staticmethod
    def load_audio(data: PathLike, language: str = "en-US") -> str:
        raise NotImplementedError(
            "Audio transcription is provided by the audio processor service."
        )

    @staticmethod
    def load_video(data: PathLike, extract_audio: bool = True) -> str:
        raise NotImplementedError("Video transcription is not implemented.")

    def load_url(self, data: str) -> str:
        import requests
        from bs4 import BeautifulSoup

        parsed_url = urlparse(data)
        if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
            raise ValueError("Only absolute HTTP and HTTPS URLs can be processed.")

        response = requests.get(
            data,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 Chrome/91.0.4472.124 Safari/537.36"
                )
            },
            timeout=self.request_timeout,
        )
        response.raise_for_status()

        soup = BeautifulSoup(response.text, "html.parser")
        for element in soup(["script", "style", "noscript"]):
            element.decompose()
        return soup.get_text(separator="\n", strip=True)

    def load_url_file(self, data: PathLike) -> str:
        url = Path(data).read_text(encoding="utf-8").strip()
        if not url:
            raise ValueError(f"URL file is empty: {data}")
        return self.load_url(url)

    @staticmethod
    def load_text(data: PathLike) -> str:
        return Path(data).read_text(encoding="utf-8")


class DataProcessor:
    def __init__(self) -> None:
        self.loader = LoadData()

    @staticmethod
    def save_result_to_txt(
        identifier: PathLike,
        result: str,
        output_dir: PathLike = "results",
    ) -> Path:
        source = Path(identifier)
        output_path = Path(output_dir)
        output_path.mkdir(parents=True, exist_ok=True)
        save_path = output_path / f"{source.stem}.txt"
        save_path.write_text(result, encoding="utf-8")
        return save_path

    def process_file(
        self,
        file_path: PathLike,
        save_to_txt: bool = True,
        **kwargs: object,
    ) -> str:
        path = Path(file_path)
        loader = self.loader.supported_formats.get(path.suffix.lower())
        if loader is None:
            raise ValueError(
                f"Unsupported file format: {path.suffix.lower() or '<none>'}"
            )

        result = loader(path, **kwargs)
        if save_to_txt:
            self.save_result_to_txt(path, result)
        return result

    def process_url(self, url: str, save_to_txt: bool = True) -> str:
        result = self.loader.load_url(url)
        if save_to_txt:
            domain = urlparse(url).netloc.replace(".", "_")
            self.save_result_to_txt(domain, result)
        return result

    def batch_process(
        self,
        folder_path: PathLike,
        save_to_txt: bool = True,
        **kwargs: object,
    ) -> Dict[str, str]:
        folder = Path(folder_path)
        if not folder.is_dir():
            raise NotADirectoryError(f"Input folder does not exist: {folder}")

        return {
            str(path): self.process_file(path, save_to_txt=save_to_txt, **kwargs)
            for path in sorted(folder.iterdir())
            if path.is_file()
        }
